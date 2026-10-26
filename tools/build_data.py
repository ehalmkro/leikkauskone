"""Builds the data file (data.js) for the Leikkauskone page from the Ministry of Finance Excel and the grouping CSVs.

Usage (from the project root):
    python3 tools/build_data.py

Reads:
    "Excel-taulukko meno- ja rakennekartoitus.xlsx"  – the Ministry of Finance (VM) list of spending measures
    data/grouping.csv                                 – categories, groups and alternatives for spending measures
    data/taxes.csv                                    – tax measures from the VM tax review (optional)
Writes:
    data.js                                           – data used by the page (do not edit by hand)

Requires openpyxl:  pip install openpyxl
"""
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
EXCEL = ROOT / "Excel-taulukko meno- ja rakennekartoitus.xlsx"
GROUPING = ROOT / "data" / "grouping.csv"
TAXES = ROOT / "data" / "taxes.csv"
TAX_CATEGORY_NAME = "Verot (verokartoitus)"  # shown in the UI
OUTPUT = ROOT / "data.js"

YEARS = list(range(2028, 2036))
FIRST_ROW = 5  # rows 1–4 are headers
# Columns (1 = A): B target, C measure, D background, E background memo,
# F–M central + local government 2028–2035, N–U social security funds 2028–2035
COL_TARGET, COL_MEASURE, COL_BACKGROUND, COL_MEMO = 2, 3, 4, 5
COL_GOVERNMENT, COL_FUNDS = 6, 14


def to_number(value):
    """Excel cell -> million euros (float), or None when there is no estimate ('-')."""
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    return None


def clean(text):
    if text is None:
        return ""
    lines = [" ".join(line.split()) for line in str(text).replace("\r", "").split("\n")]
    # keep paragraph breaks, drop repeated blank lines
    result, blank = [], False
    for line in lines:
        if line:
            result.append(line)
            blank = False
        elif not blank and result:
            result.append("")
            blank = True
    return "\n".join(result).strip()


def read_excel():
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    date = wb.worksheets[0]["A4"].value
    rows = {}
    for ws in wb.worksheets[1:]:
        for r in range(FIRST_ROW, ws.max_row + 1):
            measure = ws.cell(r, COL_MEASURE).value
            if not measure or not str(measure).strip():
                continue
            rows[(ws.title, r)] = {
                "target": clean(ws.cell(r, COL_TARGET).value),
                "measure": clean(measure),
                "background": clean(ws.cell(r, COL_BACKGROUND).value),
                "memo": clean(ws.cell(r, COL_MEMO).value),
                "government": [to_number(ws.cell(r, COL_GOVERNMENT + i).value) for i in range(8)],
                "funds": [to_number(ws.cell(r, COL_FUNDS + i).value) for i in range(8)],
            }
    return str(date or "").strip(), rows


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def csv_number(text):
    text = text.strip().replace(",", ".").replace(" ", "")
    return float(text) if text else None


def add_to_group(categories, category_id, category_name, group_name, measure_id):
    category = next((c for c in categories if c["name"] == category_name), None)
    if not category:
        category = {"id": category_id, "name": category_name, "groups": []}
        categories.append(category)
    group = next((g for g in category["groups"] if g["name"] == group_name), None)
    if not group:
        group = {"id": f"{category['id']}:{len(category['groups']) + 1}", "name": group_name, "measures": []}
        category["groups"].append(group)
    group["measures"].append(measure_id)


def read_conflicts(row, measure_id, alternative_groups, conflicts_with):
    for g in row["alternative_group"].split():
        alternative_groups.setdefault(g, []).append(measure_id)
    for other in row["conflicts_with"].split():
        conflicts_with.setdefault(measure_id, set()).add(other)
        conflicts_with.setdefault(other, set()).add(measure_id)


def main():
    date, excel = read_excel()
    grouping = read_csv(GROUPING)

    measures = {}
    categories = []  # keeps the CSV order
    alternative_groups = {}
    conflicts_with = {}
    used = set()
    warnings = []

    for row in grouping:
        mid = row["id"].strip()
        key = (row["sheet"], int(row["row"]))
        if key not in excel:
            sys.exit(f"ERROR: {mid}: row {key} not found in the Excel file")
        used.add(key)
        e = excel[key]

        name = e["measure"]
        # "TAI ..." (= "OR ...") marks an alternative to the previous measure (but not "Taiteen ...")
        alternative = re.match(r"^TAI\s", name) is not None
        if alternative:
            name = name[3:].strip()
        name = name.replace("\n", " ")

        if alternative and not row["alternative_group"].strip():
            warnings.append(f"{mid} is a TAI alternative but has no alternative_group")
        read_conflicts(row, mid, alternative_groups, conflicts_with)

        measures[mid] = {
            "id": mid,
            "type": "spending",
            "name": name,
            "target": e["target"],
            "background": e["background"],
            "memo": "" if e["memo"].lower().startswith("ei taustamuistiota") else e["memo"],
            "alternative": alternative,
            "government": e["government"],
            "funds": e["funds"],
        }

        add_to_group(categories, mid.split("-")[0], row["category"].strip(), row["group"].strip(), mid)

    # Tax review: a single yield estimate (2028 level), assumed constant for all years
    if TAXES.exists():
        for row in read_csv(TAXES):
            mid = row["id"].strip()
            value = csv_number(row["impact_meur"])
            from_year = int(row["from_year"]) if row["from_year"].strip() else YEARS[0]
            name = row["measure"].strip()
            unit = None
            if row["unit_max"].strip():
                name = re.sub(r"^YKSIKKÖTUOTTO:\s*", "", name)
                unit = {"max": int(row["unit_max"]), "step": row["unit_step"].strip()}
            read_conflicts(row, mid, alternative_groups, conflicts_with)
            measures[mid] = {
                "id": mid,
                "type": "tax",
                "name": name,
                "target": row["group"].strip(),
                "background": row["vm_comment"].strip(),
                "memo": "",
                "alternative": False,
                "government": [None if value is None else (value if y >= from_year else 0) for y in YEARS],
                "funds": [None] * len(YEARS),
                "unit": unit,
                "vmEstimate": row["vm_estimate"].strip(),
                "basis": row["basis"].strip(),
                "fromYear": from_year,
            }
            add_to_group(categories, "VER", TAX_CATEGORY_NAME, row["group"].strip(), mid)

    # Conflicts: members of the same alternative group + explicit conflicts_with pairs
    for m in measures.values():
        m["conflicts"] = []
    for members in alternative_groups.values():
        for a in members:
            measures[a]["conflicts"] += [b for b in members if b != a]
    for a, others in conflicts_with.items():
        if a not in measures:
            sys.exit(f"ERROR: conflicts_with refers to an unknown id {a}")
        measures[a]["conflicts"] += sorted(others)
    for m in measures.values():
        m["conflicts"] = sorted(set(m["conflicts"]))

    for key in sorted(set(excel) - used):
        warnings.append(f"Excel row {key} is missing from grouping.csv and will not be shown")

    data = {
        "source": {
            "name": "Valtiovarainministeriö: Julkisen talouden meno- ja rakennekartoitus, toimenpidelistaus",
            "date": date,
            "taxes": TAXES.exists(),
        },
        "years": YEARS,
        "categories": categories,
        "measures": measures,
    }
    OUTPUT.write_text(
        "// Generated by tools/build_data.py – do not edit by hand.\n"
        "window.CUTS_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    for w in warnings:
        print("WARNING:", w)
    print(f"Wrote {OUTPUT.name}: {len(measures)} measures, "
          f"{len(categories)} categories, {sum(len(c['groups']) for c in categories)} groups.")


if __name__ == "__main__":
    main()
