"""Rakentaa leikkauskoneen datatiedoston (data.js) VM:n Excelistä ja ryhmittelystä.

Käyttö (projektin juuresta):
    python3 tyokalut/rakenna_data.py

Lukee:
    "Excel-taulukko meno- ja rakennekartoitus.xlsx"  – VM:n toimenpidelistaus
    data/ryhmittely.csv                               – menojen pääluokat, alaryhmät ja vaihtoehdot
    data/verot.csv                                    – verokartoituksen toimenpiteet (valinnainen)
Kirjoittaa:
    data.js                                           – sivun käyttämä data (älä muokkaa käsin)

Vaatii openpyxl-kirjaston:  pip install openpyxl
"""
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

JUURI = Path(__file__).resolve().parent.parent
EXCEL = JUURI / "Excel-taulukko meno- ja rakennekartoitus.xlsx"
RYHMITTELY = JUURI / "data" / "ryhmittely.csv"
VEROT = JUURI / "data" / "verot.csv"
VEROT_PAALUOKKA = "Verot (verokartoitus)"
ULOS = JUURI / "data.js"

VUODET = list(range(2028, 2036))
ENSIMMAINEN_RIVI = 5  # rivit 1–4 ovat otsikoita
# Sarakkeet (1 = A): B kohde, C toimenpide, D taustaa, E taustamuistio,
# F–M valtio + paikallishallinto 2028–2035, N–U sosiaaliturvarahastot 2028–2035
SAR_KOHDE, SAR_TOIMENPIDE, SAR_TAUSTA, SAR_MUISTIO = 2, 3, 4, 5
SAR_VALTIO, SAR_SOTU = 6, 14


def luku(arvo):
    """Excelin solu -> milj. euroa (float) tai None, jos arviota ei ole ('-')."""
    if isinstance(arvo, (int, float)):
        return round(float(arvo), 2)
    return None


def siisti(teksti):
    if teksti is None:
        return ""
    rivit = [" ".join(r.split()) for r in str(teksti).replace("\r", "").split("\n")]
    # säilytetään kappalejako, poistetaan ylimääräiset tyhjät rivit
    tulos, tyhja = [], False
    for r in rivit:
        if r:
            tulos.append(r)
            tyhja = False
        elif not tyhja and tulos:
            tulos.append("")
            tyhja = True
    return "\n".join(tulos).strip()


def lue_excel():
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    paivays = wb.worksheets[0]["A4"].value
    rivit = {}
    for ws in wb.worksheets[1:]:
        for r in range(ENSIMMAINEN_RIVI, ws.max_row + 1):
            toimenpide = ws.cell(r, SAR_TOIMENPIDE).value
            if not toimenpide or not str(toimenpide).strip():
                continue
            rivit[(ws.title, r)] = {
                "kohde": siisti(ws.cell(r, SAR_KOHDE).value),
                "toimenpide": siisti(toimenpide),
                "tausta": siisti(ws.cell(r, SAR_TAUSTA).value),
                "muistio": siisti(ws.cell(r, SAR_MUISTIO).value),
                "valtio": [luku(ws.cell(r, SAR_VALTIO + i).value) for i in range(8)],
                "sotu": [luku(ws.cell(r, SAR_SOTU + i).value) for i in range(8)],
            }
    return str(paivays or "").strip(), rivit


def lue_csv(polku):
    with open(polku, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def csv_luku(teksti):
    teksti = teksti.strip().replace(",", ".").replace(" ", "")
    return float(teksti) if teksti else None


def lisaa_ryhmaan(paaluokat, pl_id, pl_nimi, ar_nimi, tid):
    pl = next((p for p in paaluokat if p["nimi"] == pl_nimi), None)
    if not pl:
        pl = {"id": pl_id, "nimi": pl_nimi, "alaryhmat": []}
        paaluokat.append(pl)
    ar = next((a for a in pl["alaryhmat"] if a["nimi"] == ar_nimi), None)
    if not ar:
        ar = {"id": f"{pl['id']}:{len(pl['alaryhmat']) + 1}", "nimi": ar_nimi, "toimenpiteet": []}
        pl["alaryhmat"].append(ar)
    ar["toimenpiteet"].append(tid)


def lue_vaihtoehdot(rivi, tid, vaihtoehtoryhmat, ei_yhdessa):
    for g in rivi["vaihtoehtoryhmä"].split():
        vaihtoehtoryhmat.setdefault(g, []).append(tid)
    for muu in rivi["ei_yhdessä"].split():
        ei_yhdessa.setdefault(tid, set()).add(muu)
        ei_yhdessa.setdefault(muu, set()).add(tid)


def main():
    paivays, excel = lue_excel()
    ryhmittely = lue_csv(RYHMITTELY)

    toimenpiteet = {}
    paaluokat = []  # säilyttää CSV:n järjestyksen
    vaihtoehtoryhmat = {}
    ei_yhdessa = {}
    kaytetyt = set()
    varoitukset = []

    for rivi in ryhmittely:
        tid = rivi["id"].strip()
        avain = (rivi["välilehti"], int(rivi["rivi"]))
        if avain not in excel:
            sys.exit(f"VIRHE: {tid}: Excelistä ei löydy riviä {avain}")
        kaytetyt.add(avain)
        e = excel[avain]

        nimi = e["toimenpide"]
        # "TAI ..." = vaihtoehto edelliselle toimelle (huom. ei "Taiteen ...")
        vaihtoehto = re.match(r"^TAI\s", nimi) is not None
        if vaihtoehto:
            nimi = nimi[3:].strip()
        nimi = nimi.replace("\n", " ")

        if vaihtoehto and not rivi["vaihtoehtoryhmä"].strip():
            varoitukset.append(f"{tid} on TAI-vaihtoehto mutta ilman vaihtoehtoryhmää")
        lue_vaihtoehdot(rivi, tid, vaihtoehtoryhmat, ei_yhdessa)

        toimenpiteet[tid] = {
            "id": tid,
            "tyyppi": "meno",
            "nimi": nimi,
            "kohde": e["kohde"],
            "tausta": e["tausta"],
            "muistio": "" if e["muistio"].lower().startswith("ei taustamuistiota") else e["muistio"],
            "vaihtoehto": vaihtoehto,
            "valtio": e["valtio"],
            "sotu": e["sotu"],
        }

        lisaa_ryhmaan(paaluokat, tid.split("-")[0], rivi["pääluokka"].strip(), rivi["alaryhmä"].strip(), tid)

    # Verokartoitus: yksi tuottoarvio (2028 tasossa), oletetaan samaksi kaikille vuosille
    if VEROT.exists():
        for rivi in lue_csv(VEROT):
            tid = rivi["id"].strip()
            arvo = csv_luku(rivi["vaikutus_milj"])
            alkaen = int(rivi["alkaen"]) if rivi["alkaen"].strip() else VUODET[0]
            nimi = rivi["toimenpide"].strip()
            yksikko = None
            if rivi["yksikkö_max"].strip():
                nimi = re.sub(r"^YKSIKKÖTUOTTO:\s*", "", nimi)
                yksikko = {"max": int(rivi["yksikkö_max"]), "askel": rivi["askel"].strip()}
            lue_vaihtoehdot(rivi, tid, vaihtoehtoryhmat, ei_yhdessa)
            toimenpiteet[tid] = {
                "id": tid,
                "tyyppi": "vero",
                "nimi": nimi,
                "kohde": rivi["alaryhmä"].strip(),
                "tausta": rivi["vm_kommentti"].strip(),
                "muistio": "",
                "vaihtoehto": False,
                "valtio": [None if arvo is None else (arvo if v >= alkaen else 0) for v in VUODET],
                "sotu": [None] * len(VUODET),
                "yksikko": yksikko,
                "vm_arvio": rivi["vm_tuottoarvio"].strip(),
                "peruste": rivi["luvun_peruste"].strip(),
                "alkaen": alkaen,
            }
            lisaa_ryhmaan(paaluokat, "VER", VEROT_PAALUOKKA, rivi["alaryhmä"].strip(), tid)

    # Ristiriidat: saman vaihtoehtoryhmän jäsenet + ei_yhdessä-parit
    for tid in toimenpiteet:
        toimenpiteet[tid]["ristiriidat"] = []
    for jasenet in vaihtoehtoryhmat.values():
        for a in jasenet:
            toimenpiteet[a]["ristiriidat"] += [b for b in jasenet if b != a]
    for a, muut in ei_yhdessa.items():
        if a not in toimenpiteet:
            sys.exit(f"VIRHE: ei_yhdessä viittaa tuntemattomaan id:hen {a}")
        toimenpiteet[a]["ristiriidat"] += sorted(muut)
    for t in toimenpiteet.values():
        t["ristiriidat"] = sorted(set(t["ristiriidat"]))

    for avain in sorted(set(excel) - kaytetyt):
        varoitukset.append(f"Excelin rivi {avain} puuttuu ryhmittelystä eikä näy koneessa")

    data = {
        "lahde": {
            "nimi": "Valtiovarainministeriö: Julkisen talouden meno- ja rakennekartoitus, toimenpidelistaus",
            "paivays": paivays,
            "verot": VEROT.exists(),
        },
        "vuodet": VUODET,
        "paaluokat": paaluokat,
        "toimenpiteet": toimenpiteet,
    }
    ULOS.write_text(
        "// Generoitu tiedostolla tyokalut/rakenna_data.py – älä muokkaa käsin.\n"
        "window.LEIKKAUSDATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    for v in varoitukset:
        print("VAROITUS:", v)
    print(f"Kirjoitettiin {ULOS.name}: {len(toimenpiteet)} toimenpidettä, "
          f"{len(paaluokat)} pääluokkaa, {sum(len(p['alaryhmat']) for p in paaluokat)} alaryhmää.")


if __name__ == "__main__":
    main()
